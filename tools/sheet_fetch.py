"""A question's row and its answer page, straight from the team's sheet, by question ID.

    python tools/sheet_fetch.py HIS-C6-SA-04 [--project his-c6-sa-04] [--refresh] [--no-vision]
    python tools/sheet_fetch.py HIS-C6-SA-04 --page <file.pdf|file.png|Drive link>   # a page someone sends later

Writes projects/<slug>/meta.json (merging: fields already set by hand are kept, the sheet's own
values go under "_sheet") and, when it finds one, the in-house answer page as assets/answer/<ID>-<n>.png
(one per page). meta.json "answer_page" records where it came from, or "pending".

WHERE THE PAGE IS
-----------------
Usually column AK, "Final Answer Photo Link", but not always, so the whole row is searched. Every link in
it is a candidate except the script and video columns, tried in this order: AK; any column whose name says
answer, photo, image or page; any other Drive link; and last column L, "Question Answer Image Link", which
is normally the TEXTBOOK SCAN. A link may be a Drive file (pdf or image) or a Drive folder of pages.

Each candidate is downloaded and checked with tools/answer_page_gate.py's own tests (portrait, ink, no
duplicate pages, and the vision read: ruled house sheet, handwriting, watermark, this question). The first
that passes is the page. One that fails is reported and skipped, so a textbook scan is never taken.

NO PAGE IS NOT A STOP. When nothing in the row passes and the EndScreenshot drop has nothing either, the
page is marked pending and the reel is rendered and delivered without it. When the page turns up (in the
sheet, or sent by someone), `tools/deliver.py --pending` or `tools/deliver.py <slug>` adds it to the
delivered video without re-rendering.

WHY THE XLSX EXPORT
-------------------
The CSV export drops hyperlinks, so AK reads as bare text; the xlsx export keeps them. Both sheets
are readable by link, so neither needs an account: the export URL and Drive's public download and
folder-listing pages are enough. The science sheets carry no AK links yet; their pages still come
from the EndScreenshot drop (inbox/answers/house/<ID>.png), which this falls back to.
"""
from __future__ import annotations

import argparse
import html
import json
import re
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "inbox" / "sheets"
# The sheets are readable by link, so their IDs are as good as the sheets themselves: they live in .env
# (SHEET_ARTS, SHEET_SCIENCE), never in this public repo.
SHEET_KEYS = {"arts": "SHEET_ARTS", "science": "SHEET_SCIENCE"}
TABS = {"HIS": ("arts", "History"), "GEO": ("arts", "Geography"), "POL": ("arts", "Political Science"),
        "ECO": ("arts", "Economics"), "SOC": ("arts", "Sociology"),
        "PHY": ("science", "Physics"), "CHE": ("science", "Chemistry"), "BIO": ("science", "Biology")}
HEADER_ROW = 3
COL = {"id": "Question ID", "subject": "Subject", "chapter": "Chapter Name (Hindi)", "category": "Question Category",
       "question": "Question", "answer": "Answer", "years": "Year", "page": "Final Answer Photo Link"}
UA = {"User-Agent": "Mozilla/5.0 (pipeline sheet_fetch)"}


def get(url: str, timeout: int = 120) -> tuple[bytes, str]:
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout) as r:
        return r.read(), r.headers.get("Content-Type", "")


def sheet_id(which: str) -> str:
    env = ROOT / ".env"
    kv = dict(l.split("=", 1) for l in env.read_text().splitlines() if "=" in l and not l.startswith("#")) if env.exists() else {}
    sid = kv.get(SHEET_KEYS[which], "").strip().strip('"')
    if not sid:
        raise SystemExit(f"{SHEET_KEYS[which]} is not set in .env: put the {which} sheet's ID there (the part of its URL after /d/)")
    return sid


def workbook(which: str, refresh: bool):
    CACHE.mkdir(parents=True, exist_ok=True)
    f = CACHE / f"{which}.xlsx"
    if refresh or not f.exists() or time.time() - f.stat().st_mtime > 3600:   # an hour old is stale enough to refetch
        data, _ = get(f"https://docs.google.com/spreadsheets/d/{sheet_id(which)}/export?format=xlsx")
        if not data.startswith(b"PK"):
            raise SystemExit(f"the {which} sheet did not export as xlsx (is it still shared by link?)")
        f.write_bytes(data)
    # data_only: the science ID column is a formula, and the export carries its computed value
    return openpyxl.load_workbook(f, data_only=True)


def find(qid: str, refresh: bool) -> dict:
    which, tab = TABS.get(qid[:3].upper(), (None, None))
    if not which:
        raise SystemExit(f"{qid}: unknown subject prefix (expected one of {', '.join(TABS)})")
    ws = workbook(which, refresh)[tab]
    head = {str(c.value or "").strip(): c.column - 1 for c in ws[HEADER_ROW]}
    # exact header first: "Question" would otherwise match "Question ID" and "Question Category" too
    idx = {k: head.get(v, next((i for h, i in head.items() if h.startswith(v)), None)) for k, v in COL.items()}
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        if str(row[idx["id"]].value or "").strip().upper() == qid.upper():
            val = lambda k: str(row[idx[k]].value or "").strip() if idx[k] is not None else ""
            names = {i: h for h, i in head.items()}
            links = []
            for c in row:
                url = c.hyperlink.target if c.hyperlink else ""
                if not url and isinstance(c.value, str) and "drive.google.com" in c.value:
                    url = re.search(r"https://\S+", c.value).group(0)
                if url:
                    links.append((names.get(c.column - 1, c.column_letter), url))
            return {"sheet": which, "tab": tab, "row": row[0].row, "id": qid.upper(), "subject": val("subject") or tab,
                    "chapter": val("chapter"), "category": val("category"), "question": val("question"),
                    "answer": val("answer"), "years": val("years"), "links": links}
    raise SystemExit(f"{qid} is not in the {tab} tab of the {which} sheet")


class FetchError(Exception):
    pass


NOT_PAGE = ("script", "heygen", "video", "manim", "veo", "teacher", "recording", "reviewer")
TEXTBOOK = "Question Answer Image Link"


def ranked(links: list[tuple[str, str]]) -> list[tuple[str, str]]:
    """Candidate answer-page links, most likely first; scripts, videos and documents are not candidates."""
    def rank(h: str, url: str) -> int | None:
        hl = h.lower()
        if any(k in hl for k in NOT_PAGE) or "docs.google.com/document" in url or "youtu" in url:
            return None
        if h.startswith(TEXTBOOK):
            return 3
        if h.startswith("Final Answer Photo"):
            return 0
        if any(k in hl for k in ("answer", "photo", "image", "page", "उत्तर")):
            return 1
        return 2 if "drive.google.com" in url else None
    scored = [(r, h, u) for h, u in links if (r := rank(h, u)) is not None]
    return [(h, u) for r, h, u in sorted(scored, key=lambda x: x[0])]


def judge(pages: list[Path], question: str, vision: bool) -> list[str]:
    """answer_page_gate's own tests on a candidate; returns the reasons it is NOT the house page."""
    sys.path.insert(0, str(ROOT / "tools"))
    import answer_page_gate as g                         # noqa: PLC0415
    found = g.cheap_checks(pages) + (g.vision_checks(pages, question) if vision else [])
    return [f"{p}: {m}" for level, p, m in found if level == "FAIL"]


def drive_id(link: str) -> tuple[str, str]:
    m = re.search(r"/folders/([\w-]+)", link)
    if m:
        return "folder", m.group(1)
    m = re.search(r"/file/d/([\w-]+)", link) or re.search(r"[?&]id=([\w-]+)", link)
    if m:
        return "file", m.group(1)
    raise FetchError(f"not a Drive link: {link}")


def fetch_page(link: str, qid: str, dest: Path, first: int = 1) -> list[Path]:
    """Pages are always <ID>-1.png, <ID>-2.png ...: tools/deliver.py and kit_preflight look for exactly that."""
    if Path(link).exists():                     # a page someone sent: a local pdf or image
        data, ctype = Path(link).read_bytes(), "local file"
        return save_page(data, ctype, qid, dest, first)
    kind, fid = drive_id(link)
    if kind == "folder":                        # a public folder lists itself on this page
        page = get(f"https://drive.google.com/embeddedfolderview?id={fid}")[0].decode("utf-8", "replace")
        entries = re.findall(r'href="https://drive.google.com/file/d/([\w-]+)[^"]*".*?flip-entry-title">([^<]+)', page, re.S)
        if not entries:
            raise FetchError("the folder lists no files (is it shared by link?)")
        pick = [e for e in entries if html.unescape(e[1]).upper().startswith(qid.upper())] or entries
        out = []
        for f, _ in sorted(pick, key=lambda e: e[1]):              # one file per page, in the folder's order
            for page in fetch_page(f"https://drive.google.com/file/d/{f}/view", qid, dest, first=len(out) + 1):
                out.append(page)
        return out
    data, ctype = get(f"https://drive.google.com/uc?export=download&id={fid}")
    return save_page(data, ctype, qid, dest, first)


def save_page(data: bytes, ctype: str, qid: str, dest: Path, first: int) -> list[Path]:
    dest.mkdir(parents=True, exist_ok=True)
    if data.startswith(b"%PDF"):
        pdf = dest / (f"{qid}.pdf" if first == 1 else f"{qid}-{first}.pdf"); pdf.write_bytes(data)
        subprocess.run(["pdftoppm", "-r", "150", "-png", str(pdf), str(dest / f"_{qid}")], check=True)
        out = []
        for i, f in enumerate(sorted(dest.glob(f"_{qid}-*.png")), first):
            im = Image.open(f).convert("RGB"); f.unlink()
            p = dest / f"{qid}-{i}.png"
            im.resize((1602, round(im.height * 1602 / im.width)), Image.LANCZOS).save(p)   # the width the outro expects
            out.append(p)
        return out
    if data[:8] == b"\x89PNG\r\n\x1a\n" or data[:3] == b"\xff\xd8\xff":
        p = dest / f"{qid}-{first}.png"; Image.open(__import__("io").BytesIO(data)).convert("RGB").save(p)
        return [p]
    raise FetchError(f"did not download as a pdf or image ({ctype or 'unknown'}); is it shared by link?")


def years_line(raw: str) -> str:
    """The card's year line: "2019, 21 और 2023" (middle years shortened). The narration's years win when they differ."""
    ys = sorted(dict.fromkeys(int(t) + (2000 if int(t) < 100 else 0) for t in re.findall(r"\d{2,4}", raw)))
    if not ys:
        return ""
    if len(ys) == 1:
        return str(ys[0])
    mid = [str(y)[2:] for y in ys[1:-1]]
    return ", ".join([str(ys[0])] + mid) + f" और {ys[-1]}"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("qid")
    ap.add_argument("--project", help="project slug (default: the ID in lower case)")
    ap.add_argument("--refresh", action="store_true", help="re-download the sheet even if the cached copy is fresh")
    ap.add_argument("--page", help="a page someone sent: a local pdf/png, or a Drive link (checked like any candidate)")
    ap.add_argument("--no-vision", action="store_true", help="skip the vision check (faster; the gate still runs it later)")
    a = ap.parse_args()
    q = find(a.qid, a.refresh)
    proj = ROOT / "projects" / (a.project or q["id"].lower()); proj.mkdir(parents=True, exist_ok=True)
    meta_p = proj / "meta.json"
    meta = json.loads(meta_p.read_text()) if meta_p.exists() else {}
    for k, v in (("subject", q["tab"]), ("sheet_id", q["id"]), ("question", q["question"]), ("category", q["category"]),
                 ("chapter", q["chapter"]), ("years", years_line(q["years"])), ("parts", 1)):
        if v and (not meta.get(k) or meta.get(k) == "pending"):
            meta[k] = v
    meta["_sheet"] = {k: q[k] for k in ("sheet", "tab", "row", "years")}
    meta_p.write_text(json.dumps(meta, ensure_ascii=False, indent=1))
    print(f"{q['id']}  {q['tab']} row {q['row']}  ->  {meta_p.relative_to(ROOT)}")
    print(f"  question  {q['question'][:90]}")
    print(f"  years     {q['years']}  (card: {meta.get('years')})")
    if "marks" not in meta:
        print("  marks     not in the sheet: set \"marks\" in meta.json from the narration (preflight fails without it)")
    dest = proj / "assets" / "answer"
    if a.page:
        cands = [("supplied", a.page)]
    else:
        cands = ranked(q["links"])
        house = ROOT / "inbox/answers/house" / f"{q['id']}.png"
        if house.exists():
            cands.append(("EndScreenshot drop", str(house)))
    got, why = None, []
    for col, url in cands:
        tmp = dest / "_candidate"
        if tmp.exists():
            for f in tmp.iterdir():
                f.unlink()
        try:
            pages = fetch_page(url, q["id"], tmp)
        except (FetchError, OSError, subprocess.CalledProcessError) as e:
            why.append(f"{col}: {e}"); continue
        bad = judge(pages, meta.get("question", ""), not a.no_vision)
        if bad:
            why.append(f"{col}: not the house page ({bad[0][:90]})"); continue
        for old in dest.glob(f"{q['id']}-*.png"):
            old.unlink()
        for f in tmp.iterdir():
            f.replace(dest / f.name)
        got = (col, sorted(dest.glob(f"{q['id']}-*.png")))
        break
    if (dest / "_candidate").exists():
        for f in (dest / "_candidate").iterdir():
            f.unlink()
        (dest / "_candidate").rmdir()
    meta["answer_page"] = got[0] if got else "pending"
    meta_p.write_text(json.dumps(meta, ensure_ascii=False, indent=1))
    for w in why:
        print(f"  skipped   {w}")
    if got:
        print(f"  answer    {len(got[1])} page(s) from {got[0]}: " + ", ".join(p.name for p in got[1] if not p.stem.endswith("_page")))
    else:
        print(f"  answer    PENDING: no house answer page in {q['id']}'s row yet. The reel renders and delivers without it;"
              f" when it is in the sheet (any column) or someone sends it, run: tools/deliver.py --pending")
    return 0


if __name__ == "__main__":
    sys.exit(main())
